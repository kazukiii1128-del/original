"""
Grosmimi インフルエンサー コンテンツガイドライン（日本版）Word生成スクリプト
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def add_cell_text(cell, text, bold=False, size=10, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ''
    para = cell.paragraphs[0]
    para.alignment = align
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_section_header(doc, text, bg_hex, text_color=(255, 255, 255), size=12):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_bg(cell, bg_hex)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(*text_color)
    doc.add_paragraph()


def add_requirements_section(doc):
    """必須要件セクションを追加"""
    # メインタイトル
    title_table = doc.add_table(rows=1, cols=1)
    title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = title_table.cell(0, 0)
    set_cell_bg(cell, '2C2C2C')
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('GROSMIMI JAPAN　インフルエンサー向けコンテンツガイドライン')
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph()

    # 前置き
    intro = doc.add_paragraph()
    ir = intro.add_run('平素よりお世話になっております。\n以下は、今回の動画制作・投稿に関する必須要件となります。ご確認の上、ご対応をお願いいたします。')
    ir.font.size = Pt(10)
    doc.add_paragraph()

    # 必須要件テーブル
    add_section_header(doc, '必須要件', 'C0392B', (255, 255, 255), 11)

    requirements = [
        ('オリジナル動画\nファイルの提供',
         '・弊社での二次利用を想定して、オリジナル動画ファイルをご提供ください。\n'
         '・使用する音楽は著作権フリーのものをご使用ください。\n'
         '・動画内に字幕（テロップ）は含めずにご提出ください。'),
        ('ホワイトリスト\nコードの付与',
         '・投稿用にホワイトリストコードへのアクセス許可をお願いいたします。'),
        ('投稿用動画\nコンテンツ',
         '・インフルエンサー様が実際にSNSにアップロードする動画です。\n'
         '・動画には必ずナレーション（音声）と字幕を含めてください。\n\n'
         '【Instagramでの投稿について】\n'
         '投稿時には共同投稿（コラボ投稿）として\nGROSMIMI（grosmimi_japan）アカウントを設定してください。'),
        ('納期',
         '・商品到着後、3週間以内にコンテンツをアップロードしてください。'),
        ('確認・承認',
         '・投稿前に弊社チームにて内容を確認・承認させていただきます。\n'
         '・承認完了後にアップロードをお願いいたします。'),
    ]

    req_table = doc.add_table(rows=len(requirements), cols=2)
    req_table.style = 'Table Grid'
    for i, (label, value) in enumerate(requirements):
        c0 = req_table.cell(i, 0)
        c1 = req_table.cell(i, 1)
        set_cell_bg(c0, 'FDEDEC')
        add_cell_text(c0, label, bold=True, size=9)
        add_cell_text(c1, value, size=9)
        c0.width = Cm(3.5)
        c1.width = Cm(13.5)
    doc.add_paragraph()
    doc.add_page_break()


def create_guideline(doc, product_type):
    is_ppsu = product_type == "PPSU"

    # ---- タイトル ----
    title_table = doc.add_table(rows=1, cols=1)
    title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = title_table.cell(0, 0)
    set_cell_bg(cell, '2C2C2C')
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('コンテンツガイドライン（日本版）')
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph()

    # ---- Campaign Overview ----
    add_section_header(doc, 'キャンペーン概要', '4A90D9', (255, 255, 255), 12)

    if is_ppsu:
        overview_data = [
            ('商品名', 'Grosmimi PPSUストローマグ'),
            ('読み方', 'グロミミ（Grosmimi）'),
            ('伝えたい課題',
             'プラスチック（PP素材）のマグは使い続けると傷がつきやすく、\n'
             'そこから微細なプラスチック粒子が溶け出す可能性がある。\n'
             '毎日赤ちゃんの口に触れるものだからこそ、素材選びが大切。'),
            ('伝えたいこと（Key Message）',
             'PPSU素材 = 医療現場でも使われる安心素材。\n'
             '傷がつきにくく、割れない、耐熱性あり、食洗機対応。\n'
             '赤ちゃんのために素材から選んだマグ。'),
            ('重要ワード', '「PPSU」「PPSUストローマグ」'),
            ('ハッシュタグ',
             '#グロミミ #Grosmimi #ストローマグ #スマートマグ #PPSU'),
            ('サブベネフィット',
             '・こぼれにくい構造\n'
             '・ストローが柔らかく、ストロー飲み練習中の赤ちゃんにも使いやすい\n'
             '・ストローが太め → お手入れしやすい\n'
             '・食洗機対応'),
        ]
    else:
        overview_data = [
            ('商品名', 'Grosmimi ステンレスストローマグ'),
            ('読み方', 'グロミミ（Grosmimi）'),
            ('伝えたいこと（Key Message）',
             '内側素材はSUS316ステンレス（食品グレード・医療現場でも使用）。\n'
             'ステンレスにもグレードの差があることを知ってほしい。\n'
             '毎日口に入れるものだから、素材の品質にこだわった。'),
            ('重要ワード', '「食品グレード ストローマグ」「SUS316」'),
            ('ハッシュタグ',
             '#グロミミ #Grosmimi #ストローマグ #スマートマグ'),
            ('サブベネフィット',
             '・完全鉛フリー（キャップにはシリコンスプリングのみ）\n'
             '・こぼれにくい構造\n'
             '・内側に目盛り付き → 飲んだ量が一目でわかる'),
        ]

    ov_table = doc.add_table(rows=len(overview_data), cols=2)
    ov_table.style = 'Table Grid'
    for i, (label, value) in enumerate(overview_data):
        cell_l = ov_table.cell(i, 0)
        cell_r = ov_table.cell(i, 1)
        set_cell_bg(cell_l, 'EFEFEF')
        add_cell_text(cell_l, label, bold=True, size=9)
        add_cell_text(cell_r, value, size=9)
        cell_l.width = Cm(4.2)
        cell_r.width = Cm(12.8)
    doc.add_paragraph()

    # ---- コンテンツ構成 ----
    cat_intro = doc.add_paragraph()
    ci = cat_intro.add_run(
        '下記のコンテンツ構成を参考に、動画の制作をお願いいたします。\n'
        'ナレーション例はあくまで参考です。ご自身の言葉で自然に話していただいて構いません。'
    )
    ci.bold = True
    ci.font.size = Pt(10)
    ci.font.color.rgb = RGBColor(30, 30, 30)
    doc.add_paragraph()

    add_section_header(doc, 'コンテンツ構成　日常 → 気づき → 共有', '4A90D9', (255, 255, 255), 12)

    subtitle = doc.add_paragraph()
    sr = subtitle.add_run('※ コンテンツの70%は「あなた自身のリアルな体験・感情」。商品の説明は30%以内に。\n'
                          '「ママ友に教えてあげる感覚」で話すのが一番伝わります。')
    sr.italic = True
    sr.font.size = Pt(9)
    sr.font.color.rgb = RGBColor(100, 100, 100)
    doc.add_paragraph()

    if is_ppsu:
        structure_data = [
            ('1. 日常フック\n（冒頭 0〜3秒）',
             '赤ちゃんが楽しそうにマグで飲んでいる自然な場面\n'
             '→ 顔と口元が両方映る角度で\n'
             '※ 商品アップ・開封シーンから入らない',
             '【パターンA】\n「ストローマグって何がいいかずっと悩んでたんやけど、\nやっと納得のいくのに出会えた」\n\n'
             '【パターンB】\n「うちの子が毎日使ってるマグ、素材のことちゃんと調べてから選んだやつで」'),

            ('2. 気づき\n（3〜12秒）',
             '使い込んで傷がついた古いプラスチックマグのアップ\n'
             '→ 新しいGrosmimiマグと並べて比較',
             '「プラスチックのマグって長く使ってると傷だらけになるじゃないですか。\nそこからマイクロプラスチックが出るって調べてわかって、\nそれからずっと気になってたんよね」'),

            ('3. 発見・解決\n（12〜25秒）',
             '赤ちゃんが飲んでいる場面（自然光・明るい背景）\n'
             '+ ストロー・本体素材のアップ\n'
             '+ 食洗機に入れてる場面があればなおよし',
             '「これ、PPSUっていう素材で作られてて、\n医療の現場でも使われてる安心素材なんよ。\n何回洗っても傷がつきにくくて、食洗機もOKで。\n赤ちゃんの口に毎日入るものやから、\n素材だけはちゃんと選びたかった」'),

            ('4. 自然な締め\n（25〜30秒）',
             '子どもが飲み終わって満足してる顔\n'
             'または自分がカメラを見て話す場面',
             '「ストローマグ選ぶときはPPSU素材かどうか\n確認してみるといいと思う。\n気になる方はチェックしてみてね」'),
        ]
    else:
        structure_data = [
            ('1. 日常フック\n（冒頭 0〜5秒）',
             '1日の中で子どもが飲んでいる場面を\n2〜3カット素早くつなぐ（朝・昼・おやつ）\n'
             '→ 日常に溶け込んでいる雰囲気を出す',
             '【パターンA】\n「うちの子、朝も昼もずっとこれで飲んでるんやけど」\n\n'
             '【パターンB】\n「ステンレスのマグって安心そうやけど、\nステンレスにもグレードがあるって知ってた？」'),

            ('2. 核心シーン\n（5〜18秒）',
             '【必須カット】\n・マグ内側底面の「SUS 316」刻印アップ\n'
             '・外側の「High quality stainless steel SUS 316」刻印アップ\n'
             '→ テキストがはっきり映るように明るい場所で撮影',
             '「ステンレスにも品質の差があって、\nこのSUS316っていうのが食品グレードの中でもトップのやつで。\n病院とか医療現場でも使われてる素材なんよ。\n内側がこの素材っていうのが決め手やった」'),

            ('3. 使ってみての感想\n（18〜28秒）',
             '赤ちゃんが自分でマグを持って飲んでいる場面\n'
             '→ 飲み終わって顔を上げた瞬間を逃さない',
             '「毎日口に入れるものやから、\nここだけは妥協したくなくて調べたら\nこれに行き着いた。\n使い始めてから素材のことを気にしなくて\nよくなったのがほんまにラクで」'),

            ('4. 自然な締め\n（28〜35秒）',
             '商品を手に持ちながら話す\nまたは子どもと一緒に映る場面',
             '「ステンレスのマグ選ぶとき、\nSUS316かどうかを確認するだけで\n全然違うと思う。\n気になった方はぜひ調べてみてね」'),
        ]

    col_headers = ['シーン', 'ビジュアル指示', 'ナレーション参考例']
    st_table = doc.add_table(rows=len(structure_data) + 1, cols=3)
    st_table.style = 'Table Grid'

    header_widths = [3.0, 6.0, 8.0]
    for j, h in enumerate(col_headers):
        cell = st_table.cell(0, j)
        set_cell_bg(cell, 'F5C518')
        add_cell_text(cell, h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell.width = Cm(header_widths[j])

    for i, (step, visual, narr) in enumerate(structure_data):
        row_cells = [st_table.cell(i + 1, j) for j in range(3)]
        set_cell_bg(row_cells[0], 'FFFDE7')
        add_cell_text(row_cells[0], step, bold=True, size=9)
        add_cell_text(row_cells[1], visual, size=9)
        add_cell_text(row_cells[2], narr, size=9)
        for j, w in enumerate(header_widths):
            row_cells[j].width = Cm(w)

    doc.add_paragraph()

    # ---- DO's / DON'Ts（日本版） ----
    add_section_header(doc, 'やってほしいこと / やらないでほしいこと', '2C2C2C', (255, 255, 255), 11)

    if is_ppsu:
        dos = [
            '1. 「ママ友に教える感覚」で話す\n商品の宣伝ではなく、自分が見つけたいいものをシェアするトーンで',
            '2. 赤ちゃんが実際に使っているシーンのアップを入れる\n顔 + 口元が両方映る角度が理想',
            '3. 自分の声でナレーションを入れ、字幕をつける\n声のトーンも「日常会話」で',
            '4. 自然光・明るい背景で撮影する\n白や淡い色の背景が映えやすい',
            '5. 表情を大切に\n商品を見せるより「使ってるときの自分や子どもの顔」が刺さる',
        ]
        donts = [
            '1. 冒頭3秒以内に商品アップ・開封シーンから入らない\nまず日常の場面から始める',
            '2. 「ブランドからいただきました」感を出さない\n自分で選んだように自然に見せる',
            '3. ウォーターマーク・テキストスタンプは入れない',
            '4. 鏡写し・反転映像はNG\n商品ラベル・ブランド名が反転して映らないよう確認',
            '5. 商品説明が長くなりすぎない\n体験談・感情が主役',
        ]
    else:
        dos = [
            '1. 「ステンレスにもグレードがある」という気づきを入れる\n知識のシェアとして伝えると自然',
            '2. SUS316の刻印アップは必須カット\n内側底面 + 外側の2箇所を明るい場所で撮る',
            '3. 自分の声でナレーションを入れ、字幕をつける\n日常会話のトーンで',
            '4. 自然光・明るい背景で撮影する',
            '5. 子どもが飲んでいる自然な場面を入れる\n演出ではなく日常の一コマを切り取る感じで',
        ]
        donts = [
            '1. 冒頭3秒以内に商品アップ・開封シーンから入らない',
            '2. 「提供してもらいました」感を出さない',
            '3. ウォーターマーク・テキストスタンプは入れない',
            '4. 鏡写し・反転映像はNG',
            '5. 商品説明が長くなりすぎない\n体験談・感情が主役',
        ]

    dd_table = doc.add_table(rows=6, cols=2)
    dd_table.style = 'Table Grid'

    do_h = dd_table.cell(0, 0)
    dont_h = dd_table.cell(0, 1)
    set_cell_bg(do_h, '27AE60')
    set_cell_bg(dont_h, 'C0392B')
    add_cell_text(do_h, 'やってほしいこと', bold=True, size=11, color=(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)
    add_cell_text(dont_h, 'やらないでほしいこと', bold=True, size=11, color=(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)

    for i in range(5):
        do_c = dd_table.cell(i + 1, 0)
        dont_c = dd_table.cell(i + 1, 1)
        bg_do = 'EAF7EF' if i % 2 == 0 else 'FFFFFF'
        bg_dont = 'FDEDEC' if i % 2 == 0 else 'FFFFFF'
        set_cell_bg(do_c, bg_do)
        set_cell_bg(dont_c, bg_dont)
        add_cell_text(do_c, dos[i], size=9)
        add_cell_text(dont_c, donts[i], size=9)
        do_c.width = Cm(8.5)
        dont_c.width = Cm(8.5)

    doc.add_paragraph()

    # ---- 撮影チェックリスト ----
    add_section_header(doc, '撮影前チェックリスト', '555555', (255, 255, 255), 11)

    checklist_items = [
        ('明るさ', '自然光推奨。室内なら窓際・昼間に撮影'),
        ('カメラ角度', '赤ちゃんの顔正面〜やや斜め前（口元と表情が両方映る角度）'),
        ('録画尺', '飲んでいる場面は10秒以上録っておく（編集でカット可）'),
        ('飲む音', 'チュッという音が入るとリアル感UP。静かな環境で撮影'),
        ('字幕', 'ナレーション全体に字幕を入れる（音なし視聴者対応）'),
        ('反転確認', 'ラベル・ブランド名が鏡写しになっていないか確認'),
        ('ウォーターマーク', '入れない'),
        ('アスペクト比', '縦型 9:16（リール・TikTok対応）で撮影推奨'),
    ]

    cl_table = doc.add_table(rows=len(checklist_items) + 1, cols=3)
    cl_table.style = 'Table Grid'
    for j, h in enumerate(['項目', '基準', '確認']):
        cell = cl_table.cell(0, j)
        set_cell_bg(cell, '555555')
        add_cell_text(cell, h, bold=True, size=9, color=(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)

    cl_widths = [3.5, 12.0, 1.5]
    for i, (item, std) in enumerate(checklist_items):
        bg = 'F5F5F5' if i % 2 == 0 else 'FFFFFF'
        for j in range(3):
            c = cl_table.cell(i + 1, j)
            set_cell_bg(c, bg)
            c.width = Cm(cl_widths[j])
        add_cell_text(cl_table.cell(i + 1, 0), item, bold=True, size=9)
        add_cell_text(cl_table.cell(i + 1, 1), std, size=9)
        add_cell_text(cl_table.cell(i + 1, 2), '□', size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph()

    # ---- NG事項 ----
    add_section_header(doc, 'NG事項（必ずご確認ください）', 'C0392B', (255, 255, 255), 11)

    ng_items = [
        ('冒頭の商品アップ',
         '動画の冒頭3秒以内に商品アップ・開封シーンから入らないでください。\n'
         'まず日常シーン・体験談から始めてください。'),
        ('宣伝・PR感',
         '「ブランドからいただきました」「提供品です」という表現や雰囲気を出さないでください。\n'
         '自分で選んだアイテムとして自然に紹介してください。'),
        ('ウォーターマーク',
         'テキストスタンプ・ウォーターマーク・他社ロゴは動画内に入れないでください。'),
        ('鏡写し・反転映像',
         '商品ラベル・ブランド名（Grosmimi）が反転して映る鏡写し映像はNGです。\n'
         '投稿前に文字が正しく読めるか必ず確認してください。'),
        ('著作権のある音楽',
         '著作権フリーの音楽のみ使用してください。\n'
         '人気楽曲・テレビ曲・アーティスト楽曲はNGです（二次利用時に問題になります）。'),
        ('未承認での投稿',
         '動画の投稿前に必ず弊社チームの確認・承認を受けてください。\n'
         '承認前のアップロードはお控えください。'),
    ]

    ng_table = doc.add_table(rows=len(ng_items), cols=2)
    ng_table.style = 'Table Grid'
    for i, (label, value) in enumerate(ng_items):
        c0 = ng_table.cell(i, 0)
        c1 = ng_table.cell(i, 1)
        set_cell_bg(c0, 'FDEDEC')
        set_cell_bg(c1, 'FFF5F5' if i % 2 == 0 else 'FFFFFF')
        add_cell_text(c0, label, bold=True, size=9, color=(180, 0, 0))
        add_cell_text(c1, value, size=9)
        c0.width = Cm(3.5)
        c1.width = Cm(13.5)

    doc.add_paragraph()

    # ---- タグ ----
    tag_table = doc.add_table(rows=3, cols=2)
    tag_table.style = 'Table Grid'
    set_cell_bg(tag_table.cell(0, 0), 'C0392B')
    set_cell_bg(tag_table.cell(0, 1), 'C0392B')
    add_cell_text(tag_table.cell(0, 0), 'TikTok', bold=True, size=9, color=(255, 255, 255))
    add_cell_text(tag_table.cell(0, 1), '@grosmimi_japan', size=9, color=(255, 255, 255))
    add_cell_text(tag_table.cell(1, 0), 'Instagram', bold=True, size=9)
    add_cell_text(tag_table.cell(1, 1), '@grosmimi_japan', size=9)
    add_cell_text(tag_table.cell(2, 0), '投稿後', bold=True, size=9)
    add_cell_text(tag_table.cell(2, 1), '投稿URLをDMで共有してください', size=9)
    for row in tag_table.rows:
        row.cells[0].width = Cm(3.0)
        row.cells[1].width = Cm(14.0)


def main():
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)

    doc.styles['Normal'].font.name = 'Yu Gothic'
    doc.styles['Normal'].font.size = Pt(10)

    add_requirements_section(doc)
    create_guideline(doc, "PPSU")
    doc.add_page_break()
    create_guideline(doc, "Stainless")

    output_path = ".tmp/Grosmimi_Content_Guideline_JP.docx"
    doc.save(output_path)
    print(f"Done: {output_path}")


if __name__ == "__main__":
    main()
